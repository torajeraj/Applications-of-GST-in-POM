# grey_mrp_simple.py
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import json
from collections import defaultdict, deque
from math import ceil

# Excel handling
try:
    from openpyxl import Workbook, load_workbook
except Exception:
    Workbook = None
    load_workbook = None

# Reuse existing comparator for discrete grey numbers
from grey_comparison import DiscreteGreyNumberComparator


class GreyMRPSimpleInput(tk.Frame):
    """
    Grey MRP — Simple Approach (corrected)
    - Input: single Excel file with 3 sheets: Items, BOM, MPS & Costs
    - Buttons: Download template, Browse, Preview, Run, Save/Load
    - For each item and each candidate LT_order in [LT_low, LT_up]:
        * Builds nominal POW and PR_nom (on-time receipts when r == LT_order)
        * Evaluates all realized LT scenarios r in the same set by shifting PR_nom
        * Computes per-week inventory and cost (holding/shortage) EVERY week using target SS
        * Aggregates a discrete grey cost set {C(r)} for the policy
      Then compares policies using DiscreteGreyNumberComparator (on NEGATED costs)
      and selects the best policy; supports the single-policy (non-grey LT) short-circuit.
    """
    def __init__(self, parent, result_widget):
        super().__init__(parent, bg="#2a2a2a")
        self.result_widget = result_widget

        self.file_var = tk.StringVar(value="")

        # ===== Top: file chooser and template =====
        top = tk.Frame(self, bg="#2a2a2a"); top.pack(fill="x", pady=6)
        tk.Label(top, text="Excel file (Items/BOM/MPS):", bg="#2a2a2a", fg="white").pack(side="left")
        tk.Entry(top, textvariable=self.file_var, width=60).pack(side="left", padx=6)
        tk.Button(top, text="Browse", command=self.browse_file).pack(side="left", padx=4)
        tk.Button(top, text="Download template", command=self.download_template).pack(side="left", padx=8)

        # ===== Footer buttons =====
        foot = tk.Frame(self, bg="#2a2a2a"); foot.pack(fill="x", pady=8)
        tk.Button(foot, text="Save Project", command=self.save_project).pack(side="left", padx=5)
        tk.Button(foot, text="Load Project", command=self.load_project).pack(side="left", padx=5)
        tk.Button(foot, text="Preview", command=self.preview).pack(side="left", padx=5)
        tk.Button(foot, text="Run", command=self.run).pack(side="left", padx=5)

        # caches
        self.model = None     # {'items':{}, 'bom':[(p,c,u),...], 'mps':{it:{h,p,gr}}, 'weeks':[labels]}
        self.results = None

    # ---------------------- File & Template ----------------------
    def browse_file(self):
        path = filedialog.askopenfilename(filetypes=[("Excel files","*.xlsx *.xlsm *.xltx *.xltm")])
        if not path:
            return
        self.file_var.set(path)
        self.model = None
        self.results = None
        self.result_widget.delete("1.0", tk.END)
        self.result_widget.insert(tk.END, f"Selected file:\n{path}\n")

    def download_template(self):
        if Workbook is None:
            messagebox.showerror("Missing dependency", "openpyxl is required to create the Excel template.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel","*.xlsx")], title="Save Excel template as")
        if not path:
            return
        wb = Workbook()

        # Sheet 1: Items
        ws1 = wb.active
        ws1.title = "Items"
        ws1.append(["Item name","Lower Leadtime","upper Leadtime","Lot quantity (0=L4L)","safety stock","Inventory (initial)"])
        ws1.append(["A", 1, 4, 0, 0, 0])
        ws1.append(["H", 1, 3, 20, 0, 5])
        ws1.append(["F", 2, 3, 5,  0, 3])

        # Sheet 2: BOM
        ws2 = wb.create_sheet("BOM")
        ws2.append(["parent","child","usage"])
        ws2.append(["A","H",3])
        ws2.append(["A","F",2])
        ws2.append(["H","F",5])

        # Sheet 3: MPS & Costs
        ws3 = wb.create_sheet("MPS & Costs")
        # Example weeks (W13..W18)
        header = ["Item name","Holding cost per w","Slack cost per w"] + [f"W{i}" for i in range(13,19)]
        ws3.append(header)
        ws3.append(["A",10,100,0,0,0,60,70,50])
        ws3.append(["H",10,100] + [0]*6)
        ws3.append(["F",10,100] + [0]*6)

        try:
            wb.save(path)
        except Exception as e:
            messagebox.showerror("Error", f"Failed to save template:\n{e}")
            return
        messagebox.showinfo("Template created", f"Excel template saved to:\n{path}")

    # ---------------------- Preview & Run ----------------------
    def _ensure_loaded(self):
        if self.model is not None:
            return True
        path = self.file_var.get().strip()
        if not path:
            messagebox.showwarning("No file","Please choose an Excel file first.")
            return False
        if load_workbook is None:
            messagebox.showerror("Missing dependency", "openpyxl is required to read Excel files.")
            return False
        try:
            self.model = self._read_excel_model(path)
            return True
        except Exception as e:
            self.model = None
            messagebox.showerror("Read error", str(e))
            return False

    def preview(self):
        if not self._ensure_loaded():
            return
        items = self.model['items']
        bom = self.model['bom']
        mps = self.model['mps']
        weeks = self.model['weeks']

        self.result_widget.delete("1.0", tk.END)
        self.result_widget.insert(tk.END, "Grey MRP — Simple Approach (Preview)\n\n")
        self.result_widget.insert(tk.END, f"Weeks: {weeks}\n\n")

        self.result_widget.insert(tk.END, "Items:\n")
        for it, rec in items.items():
            self.result_widget.insert(tk.END, f"  - {it}: LT=[{rec['lt_low']},{rec['lt_up']}], Lot={rec['lot']}, SS={rec['ss']}, I0={rec['I0']}\n")

        self.result_widget.insert(tk.END, "\nBOM (parent -> child x usage):\n")
        if not bom:
            self.result_widget.insert(tk.END, "  (empty)\n")
        else:
            for (p,c,u) in bom:
                self.result_widget.insert(tk.END, f"  - {p} -> {c} x {u}\n")

        self.result_widget.insert(tk.END, "\nMPS & Costs (per item):\n")
        for it, row in mps.items():
            self.result_widget.insert(tk.END, f"  - {it}: h={row['h']}, p={row['p']}, MPS={row['gr']}\n")

    def run(self):
        if not self._ensure_loaded():
            return
        items = self.model['items']
        bom = self.model['bom']
        mps = self.model['mps']
        weeks = self.model['weeks']
        H = len(weeks)

        # Topological order (parents before children)
        order = self._topo_order(items.keys(), bom)

        # Store final results
        POW = {it: [0.0]*H for it in items}
        lt_choice = {}
        logs = []

        for it in order:
            rec = items[it]
            lt_low, lt_up = int(rec['lt_low']), int(rec['lt_up'])
            lot = float(rec['lot'])
            ss = float(rec['ss'])
            I0 = float(rec['I0'])

            # Build gross requirements GR: MPS (if any) + parents' POW * usage
            GR = mps[it]['gr'][:] if it in mps else [0.0]*H
            for (p, c, u) in bom:
                if c == it:
                    parent_pow = POW[p]
                    for t in range(H):
                        GR[t] += u * parent_pow[t]

            policies = list(range(lt_low, lt_up+1))
            item_logs = [f"\n{'='*72}\nITEM {it}  |  Candidate LT_order set = {policies}"]
            policy_cost_sets = []
            policy_pow_cache = {}

            for LT in policies:
                # --- Nominal plan for policy LT ---
                pow_vec = [0.0]*H
                pr_nom = [0.0]*H  # arrivals when realized r == LT (on-time)
                I_prev = I0
                item_logs.append(f"\n-- Policy LT_order = {LT} --")
                item_logs.append("Week | GR | PR_nom | Avail | NR | OQ | Rel_wk | EOH")

                for t in range(H):
                    avail = I_prev + pr_nom[t]  # include on-time receipts
                    # Safety stock target
                    NR = max(0.0, GR[t] + ss - avail)

                    if NR > 0:
                        OQ = NR if lot == 0 else ceil(NR / lot) * lot
                        rel_week = t - LT  # release timing
                        if 0 <= rel_week < H:
                            pow_vec[rel_week] += OQ
                        pr_nom[t] += OQ      # arrives on time in nominal plan
                        avail += OQ
                    else:
                        OQ = 0.0
                        rel_week = "-"

                    EOH = avail - GR[t]
                    item_logs.append(f"{t+1:>4} | {GR[t]:>6.2f} | {pr_nom[t]:>6.2f} | {avail:>6.2f} | {NR:>6.2f} | {OQ:>6.2f} | {rel_week:>6} | {EOH:>6.2f}")
                    I_prev = EOH

                item_logs.append(f"POW (release plan) for LT={LT}: {pow_vec}")
                item_logs.append(f"PR_nominal (r=LT) for LT={LT}: {pr_nom}")
                policy_pow_cache[LT] = pow_vec[:]

                # --- Evaluate realized lead-time scenarios r ---
                costs = []
                for r in policies:
                    delta = r - LT
                    # Shift nominal receipts
                    PR = [0.0]*H
                    for t in range(H):
                        q = pr_nom[t]
                        aw = t + delta
                        if 0 <= aw < H:
                            PR[aw] += q

                    I = I0
                    hold = 0.0
                    slack = 0.0
                    item_logs.append(f"   Scenario r={r}:")
                    item_logs.append("   Week | I_prev | PR | GR | I_end | cost_tag")
                    for t in range(H):
                        PR_t = PR[t]
                        GR_t = GR[t]
                        I_end = I + PR_t - GR_t

                        # ---- CORRECTED COST: compute EVERY week using target SS ----
                        cost_tag = ""
                        h = mps[it]['h'] if it in mps else 0.0
                        p = mps[it]['p'] if it in mps else 0.0
                        I_eff = I_end - ss
                        if I_eff > 0:
                            c = I_eff * h
                            hold += c
                            cost_tag = f"H={c:.2f}"
                        elif I_eff < 0:
                            c = (-I_eff) * p
                            slack += c
                            cost_tag = f"S={c:.2f}"

                        item_logs.append(f"   {t+1:>4} | {I:>6.2f} | {PR_t:>6.2f} | {GR_t:>6.2f} | {I_end:>6.2f} | {cost_tag}")
                        I = I_end

                    C = hold + slack
                    costs.append(C)
                    item_logs.append(f"   -> Cost C(LT={LT}, r={r}) = {C:.2f}  (hold={hold:.2f}, slack={slack:.2f})")

                item_logs.append(f"Grey cost set for LT={LT}: {costs}")
                policy_cost_sets.append([-c for c in costs])  # NEGATE: greater is better in comparator

            # --- Selection ---
            if len(policies) == 1:
                chosen_LT = policies[0]
                lt_choice[it] = chosen_LT
                POW[it] = policy_pow_cache[chosen_LT]
                logs.extend(item_logs)
                logs.append("\nOnly one policy available; skipping grey comparison.")
                logs.append(f"Chosen policy for {it}: LT={chosen_LT}  → POW={POW[it]}")
                continue

            comp_items = [(f"LT={LT}", ",".join(str(v) for v in vals), "") for LT, vals in zip(policies, policy_cost_sets)]
            try:
                labels, M_gt, M_eq, M_adv, scores, order_desc, log_text = DiscreteGreyNumberComparator.compare_many(comp_items)
            except Exception as e:
                messagebox.showerror("Comparison error", f"Grey discrete comparison failed for item {it}:\n{e}")
                return

            best_idx = order_desc[0]
            best_label = labels[best_idx]
            try:
                chosen_LT = int(best_label.split("=")[1])
            except Exception:
                chosen_LT = policies[0]

            lt_choice[it] = chosen_LT
            POW[it] = policy_pow_cache[chosen_LT]

            logs.extend(item_logs)
            logs.append("\nPairwise policy comparison (negated costs):")
            logs.append(log_text)
            logs.append(f"\nChosen policy for {it}: {best_label}  → POW={POW[it]}")

        # Final POW table
        logs.append("\n" + "="*72)
        logs.append("FINAL Planned Order Releases (POW):")
        head = "Item \\ Week  " + "  ".join(f"{w:>6}" for w in weeks)
        logs.append(head)
        for it in order:
            row = f"{it:>11}  " + "  ".join(f"{x:>6.2f}" for x in POW[it])
            logs.append(row)

        self.results = {"pow": POW, "lt_choice": lt_choice, "order": order, "weeks": weeks}
        self.result_widget.delete("1.0", tk.END)
        self.result_widget.insert(tk.END, "\n".join(logs) + "\n")

    # ---------------------- Excel reading helpers ----------------------
    def _read_excel_model(self, path):
        wb = load_workbook(path, data_only=True)

        # locate sheets by name fragments
        def find_sheet(names):
            for s in wb.sheetnames:
                low = s.strip().lower()
                for k in names:
                    if k in low:
                        return wb[s]
            return None

        ws_items = find_sheet(["item"])
        ws_bom = find_sheet(["bom"])
        ws_mps = find_sheet(["mps"]) or find_sheet(["mps & costs","mps&costs","cost"])

        if ws_items is None or ws_bom is None or ws_mps is None:
            raise ValueError("Cannot find all required sheets. Need: Items, BOM, MPS & Costs.")

        # Items
        items = {}
        rows = list(ws_items.iter_rows(values_only=True))
        if not rows:
            raise ValueError("Items sheet is empty.")
        header = [str(x).strip().lower() if x is not None else "" for x in rows[0]]
        def idx(colname):
            for i, h in enumerate(header):
                if colname in h:
                    return i
            return None
        i_item = idx("item")
        i_ll = idx("lower")
        i_ul = idx("upper")
        i_lot = idx("lot")
        i_ss = idx("safety")
        i_inv = idx("invent")
        for r in rows[1:]:
            if r is None: continue
            name = (r[i_item] or "").strip() if (i_item is not None and r[i_item] is not None) else ""
            if not name: continue
            items[name] = {
                "lt_low": int(r[i_ll]) if (i_ll is not None and r[i_ll] is not None) else 1,
                "lt_up":  int(r[i_ul]) if (i_ul is not None and r[i_ul] is not None) else 1,
                "lot": float(r[i_lot]) if (i_lot is not None and r[i_lot] is not None) else 0.0,
                "ss":  float(r[i_ss])  if (i_ss  is not None and r[i_ss]  is not None) else 0.0,
                "I0":  float(r[i_inv]) if (i_inv is not None and r[i_inv] is not None) else 0.0,
            }

        # BOM
        bom = []
        rows = list(ws_bom.iter_rows(values_only=True))
        if rows:
            h = [str(x).strip().lower() if x is not None else "" for x in rows[0]]
            i_p = next((i for i,v in enumerate(h) if "parent" in v), 0)
            i_c = next((i for i,v in enumerate(h) if "child" in v), 1)
            i_u = next((i for i,v in enumerate(h) if "usage" in v), 2)
            for r in rows[1:]:
                if r is None: continue
                p = (r[i_p] or "").strip() if r[i_p] is not None else ""
                c = (r[i_c] or "").strip() if r[i_c] is not None else ""
                if not p or not c: continue
                u = float(r[i_u]) if r[i_u] is not None else 1.0
                bom.append((p,c,u))

        # MPS & Costs
        mps = {}
        rows = list(ws_mps.iter_rows(values_only=True))
        if not rows:
            raise ValueError("MPS & Costs sheet is empty.")
        header = [str(x).strip() if x is not None else "" for x in rows[0]]
        # week columns
        week_cols = []
        for j,h in enumerate(header):
            if h and h.strip().lower().startswith("w"):
                week_cols.append(j)
        if not week_cols:
            raise ValueError("No week columns (W1, W2, ...) found in MPS & Costs.")
        weeks = [header[j] for j in week_cols]
        i_item = next((i for i,v in enumerate(header) if "item" in v.lower()), None)
        i_h = next((i for i,v in enumerate(header) if "holding" in v.lower()), None)
        i_p = next((i for i,v in enumerate(header) if "slack" in v.lower() or "shortage" in v.lower() or "penalty" in v.lower()), None)

        for r in rows[1:]:
            if r is None: continue
            name = (r[i_item] or "").strip() if (i_item is not None and r[i_item] is not None) else ""
            if not name: continue
            h = float(r[i_h]) if (i_h is not None and r[i_h] is not None) else 0.0
            p = float(r[i_p]) if (i_p is not None and r[i_p] is not None) else 0.0
            gr = [float(r[j] or 0.0) for j in week_cols]
            mps[name] = {"h":h, "p":p, "gr":gr}

        return {"items":items, "bom":bom, "mps":mps, "weeks":weeks}

    # ---------------------- Utils ----------------------
    def _topo_order(self, items, bom):
        items = list(items)
        indeg = {x:0 for x in items}
        children = defaultdict(list)
        for p,c,u in bom:
            if p not in indeg: indeg[p]=0
            if c not in indeg: indeg[c]=0
            indeg[c] += 1
            children[p].append(c)
        q = deque([x for x in indeg if indeg[x]==0])
        order = []
        while q:
            u = q.popleft()
            if u in order:
                continue
            if u in items:
                order.append(u)
            for v in children.get(u, []):
                indeg[v] -= 1
                if indeg[v]==0:
                    q.append(v)
        for x in items:
            if x not in order:
                order.append(x)
        return order

    # ---------------------- Save/Load project ----------------------
    def save_project(self):
        data = {"excel_path": self.file_var.get().strip()}
        path = filedialog.asksaveasfilename(defaultextension=".json", filetypes=[("JSON","*.json")], title="Save Grey MRP Simple Project")
        if not path: return
        try:
            with open(path, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=2, ensure_ascii=False)
            messagebox.showinfo("Saved", f"Saved to:\n{path}")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to save:\n{e}")

    def load_project(self):
        path = filedialog.askopenfilename(defaultextension=".json", filetypes=[("JSON","*.json")], title="Load Grey MRP Simple Project")
        if not path: return
        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
        except Exception as e:
            messagebox.showerror("Error", f"Failed to load:\n{e}")
            return
        self.file_var.set(data.get("excel_path",""))
        self.model = None
        self.results = None
        self.result_widget.delete("1.0", tk.END)
        self.result_widget.insert(tk.END, "Project loaded. Use Preview or Run.\n")

