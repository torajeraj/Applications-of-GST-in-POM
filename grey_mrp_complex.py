# grey_mrp_complex.py
# -*- coding: utf-8 -*-
import tkinter as tk
from tkinter import messagebox, filedialog
import json
from collections import defaultdict, deque
from math import ceil
from itertools import product

# Excel handling
try:
    from openpyxl import Workbook, load_workbook
except Exception:
    Workbook = None
    load_workbook = None

# Reuse the discrete grey number comparator (already in your project)
from grey_comparison import DiscreteGreyNumberComparator


class GreyMRPComplexInput(tk.Frame):
    """
    Grey MRP — Complex Approach
    UI and input format are identical to the Simple approach.
    Differences:
      - Policies are vectors of LT_order over ONLY demand weeks (GR>0).
      - For each policy, we enumerate ALL realized lead-time vectors r over the same weeks.
      - For each scenario, we simulate week-by-week and compute costs EVERY week vs SS.
      - We aggregate scenario costs into a discrete grey set (list), NEGATE for comparison,
        and select the best policy using DiscreteGreyNumberComparator.
    Logging:
      - FULL step-by-step logs (like Simple) for each policy and each scenario.
      - No pairwise comparison matrices; we print ONLY final grey scores + chosen policy.
      - Child GR build is printed (parent POW * usage) to verify alignment (e.g., A/H -> F).
    """
    def __init__(self, parent, result_widget):
        super().__init__(parent, bg="#2a2a2a")
        self.result_widget = result_widget

        self.file_var = tk.StringVar(value="")

        # ===== Top: file chooser and template =====
        top = tk.Frame(self, bg="#2a2a2a"); top.pack(fill="x", pady=6)
        tk.Label(top, text="Excel file (Items/BOM/MPS):", bg="#2a2a2a", fg="white").pack(side="left")
        tk.Entry(top, textvariable=self.file_var, width=60).pack(side="left", padx=6)
        tk.Button(top, text="Browse", command=self.browse_file).pack(side="left", padx=4)
        tk.Button(top, text="Download template", command=self.download_template).pack(side="left", padx=8)

        # ===== Footer buttons =====
        foot = tk.Frame(self, bg="#2a2a2a"); foot.pack(fill="x", pady=8)
        tk.Button(foot, text="Save Project", command=self.save_project).pack(side="left", padx=5)
        tk.Button(foot, text="Load Project", command=self.load_project).pack(side="left", padx=5)
        tk.Button(foot, text="Preview", command=self.preview).pack(side="left", padx=5)
        tk.Button(foot, text="Run", command=self.run).pack(side="left", padx=5)

        # caches
        self.model = None
        self.results = None

    # ---------------------- File & Template ----------------------
    def browse_file(self):
        path = filedialog.askopenfilename(filetypes=[("Excel files","*.xlsx *.xlsm *.xltx *.xltm")])
        if not path:
            return
        self.file_var.set(path)
        self.model = None
        self.results = None
        self.result_widget.delete("1.0", tk.END)
        self.result_widget.insert(tk.END, f"Selected file:\n{path}\n")

    def download_template(self):
        if Workbook is None:
            messagebox.showerror("Missing dependency", "openpyxl is required to create the Excel template.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel","*.xlsx")], title="Save Excel template as")
        if not path:
            return
        wb = Workbook()

        # Sheet 1: Items
        ws1 = wb.active
        ws1.title = "Items"
        ws1.append(["Item name","Lower Leadtime","upper Leadtime","Lot quantity (0=L4L)","safety stock","Inventory (initial)"])
        ws1.append(["A", 1, 4, 0, 0, 0])
        ws1.append(["H", 1, 3, 20, 0, 5])
        ws1.append(["F", 2, 3, 5,  0, 3])

        # Sheet 2: BOM
        ws2 = wb.create_sheet("BOM")
        ws2.append(["parent","child","usage"])
        ws2.append(["A","H",3])
        ws2.append(["A","F",2])
        ws2.append(["H","F",5])

        # Sheet 3: MPS & Costs
        ws3 = wb.create_sheet("MPS & Costs")
        header = ["Item name","Holding cost per w","Slack cost per w"] + [f"W{i}" for i in range(9,11)]
        ws3.append(header)
        ws3.append(["A",10,100,60,70])  # W9=60, W10=70
        ws3.append(["H",10,100,0,0])
        ws3.append(["F",10,100,0,0])

        try:
            wb.save(path)
        except Exception as e:
            messagebox.showerror("Error", f"Failed to save template:\n{e}")
            return
        messagebox.showinfo("Template created", f"Excel template saved to:\n{path}")

    # ---------------------- Preview & Run ----------------------
    def _ensure_loaded(self):
        if self.model is not None:
            return True
        path = self.file_var.get().strip()
        if not path:
            messagebox.showwarning("No file","Please choose an Excel file first.")
            return False
        if load_workbook is None:
            messagebox.showerror("Missing dependency", "openpyxl is required to read Excel files.")
            return False
        try:
            self.model = self._read_excel_model(path)
            return True
        except Exception as e:
            self.model = None
            messagebox.showerror("Read error", str(e))
            return False

    def preview(self):
        if not self._ensure_loaded():
            return
        items = self.model['items']; bom = self.model['bom']; mps = self.model['mps']; weeks = self.model['weeks']

        self.result_widget.delete("1.0", tk.END)
        self.result_widget.insert(tk.END, "Grey MRP — Complex Approach (Preview)\n\n")
        self.result_widget.insert(tk.END, f"Weeks: {weeks}\n\n")

        self.result_widget.insert(tk.END, "Items:\n")
        for it, rec in items.items():
            self.result_widget.insert(tk.END, f"  - {it}: LT=[{rec['lt_low']},{rec['lt_up']}], Lot={rec['lot']}, SS={rec['ss']}, I0={rec['I0']}\n")

        self.result_widget.insert(tk.END, "\nBOM (parent -> child x usage):\n")
        if not bom:
            self.result_widget.insert(tk.END, "  (empty)\n")
        else:
            for (p,c,u) in bom:
                self.result_widget.insert(tk.END, f"  - {p} -> {c} x {u}\n")

        self.result_widget.insert(tk.END, "\nMPS & Costs (per item):\n")
        for it, row in mps.items():
            self.result_widget.insert(tk.END, f"  - {it}: h={row['h']}, p={row['p']}, MPS={row['gr']}\n")

    def run(self):
        if not self._ensure_loaded():
            return
        items = self.model['items']; bom = self.model['bom']; mps = self.model['mps']; weeks = self.model['weeks']
        H = len(weeks)

        order = self._topo_order(items.keys(), bom)
        POW = {it: [0.0]*H for it in items}
        PRNOM = {it: [0.0]*H for it in items}   # store PR_nominal of chosen policy for each item
        lt_policy_choice = {}   # per item: dict week->LT_order_t (for weeks with GR>0)
        logs = []

        # Helper to show how child GR is built from parents' POW
        def log_child_gr_build(it, GR, bom, POW, weeks):
            lines = []
            parents = [(p, u) for (p,c,u) in bom if c == it]
            if not parents:
                return []
            lines.append("  Child GR build (sum over parents POW * usage):")
            header = "      Week  " + "  ".join(f"{w:>6}" for w in weeks)
            lines.append(header)
            # contributions per parent
            contrib_total = [0.0]*len(weeks)
            for (p, u) in parents:
                row = [u * POW[p][t] for t in range(len(weeks))]
                contrib_total = [contrib_total[t] + row[t] for t in range(len(weeks))]
                lines.append(f"      from {p} (u={u}): " + "  ".join(f"{x:>6.2f}" for x in row))
            lines.append("      TOTAL parent contribs: " + "  ".join(f"{x:>6.2f}" for x in contrib_total))
            lines.append("      MPS self GR:           " + "  ".join(f"{x:>6.2f}" for x in (mps[it]['gr'] if it in mps else [0.0]*len(weeks))))
            lines.append("      Final GR used:         " + "  ".join(f"{x:>6.2f}" for x in GR))
            return lines

        for it in order:
            rec = items[it]
            L_low, L_up = int(rec['lt_low']), int(rec['lt_up'])
            LTS = list(range(L_low, L_up+1))
            lot = float(rec['lot']); ss = float(rec['ss']); I0 = float(rec['I0'])

            # Build GR: own MPS + parents' POW*usage
            GR = self._build_gr(it, H, mps, POW, bom)
            logs.append(f"\n{'='*72}\nITEM {it}")
            logs.extend(log_child_gr_build(it, GR, bom, POW, weeks))

            # Active demand weeks (to limit Cartesian explosion)
            demand_weeks = [t for t in range(H) if GR[t] > 0]
            if not demand_weeks:
                logs.append(f"  No demand weeks; POW stays zero.")
                lt_policy_choice[it] = {}
                continue

            logs.append(f"  Complex policy over demand weeks { [w+1 for w in demand_weeks] }; LT set={LTS}")
            policies = list(product(LTS, repeat=len(demand_weeks)))
            logs.append(f"  Total policies to evaluate: {len(policies)}")

            # Cost rates
            if it not in mps:
                messagebox.showerror("Missing costs", f"No MPS & Costs row for item {it}. Please provide h and p.")
                return
            h_cost = float(mps[it]['h']); p_cost = float(mps[it]['p'])

            # ------- Helpers that log like Simple -------
            def nominal_for_policy(policy_tuple, log_list):
                # Returns POW vector and PR_on_time vector, and prints a Simple-like table
                POW_vec = [0.0]*H
                PR_nom = [0.0]*H
                I_prev = I0
                log_list.append(f"\n-- Policy LT_orders (per demand week): {policy_tuple} --")
                log_list.append("Week | GR | PR_nom | Avail | NR | OQ | Rel_wk | EOH")
                lt_map = {dw: policy_tuple[i] for i, dw in enumerate(demand_weeks)}
                for t in range(H):
                    avail = I_prev + PR_nom[t]
                    if t in lt_map:
                        NR = max(0.0, GR[t] + ss - avail)
                        if NR > 0:
                            OQ = NR if lot == 0 else ceil(NR/lot) * lot
                            rel_week = t - lt_map[t]
                            if 0 <= rel_week < H:
                                POW_vec[rel_week] += OQ
                            # Regardless of rel_week domain, receipt is at t (nominal)
                            PR_nom[t] += OQ
                            avail += OQ
                        else:
                            OQ = 0.0
                            rel_week = "-"
                    else:
                        NR = 0.0
                        OQ = 0.0
                        rel_week = "-"
                    EOH = avail - GR[t]
                    log_list.append(f"{t+1:>4} | {GR[t]:>6.2f} | {PR_nom[t]:>6.2f} | {avail:>6.2f} | {NR:>6.2f} | {OQ:>6.2f} | {str(rel_week):>6} | {EOH:>6.2f}")
                    I_prev = EOH
                log_list.append(f"POW (release plan) for policy: {POW_vec}")
                log_list.append(f"PR_nominal (on-time) for policy: {PR_nom}")
                return POW_vec, PR_nom

            def cost_for_policy(policy_tuple, log_list):
                POW_vec, PR_nom = nominal_for_policy(policy_tuple, log_list)
                cost_set = []
                for ridx, realized in enumerate(product(LTS, repeat=len(demand_weeks)), start=1):
                    PR = [0.0]*H
                    for idx, t in enumerate(demand_weeks):
                        LT_t = policy_tuple[idx]; r_t = realized[idx]
                        delta = r_t - LT_t
                        aw = t + delta
                        q = PR_nom[t]
                        if 0 <= aw < H:
                            PR[aw] += q
                    I = I0; hold = 0.0; slack = 0.0
                    log_list.append(f"   Scenario #{ridx}  r_t per demand week = {realized}")
                    log_list.append("   Week | I_prev | PR | GR | I_end | cost_tag")
                    for W in range(H):
                        I_end = I + PR[W] - GR[W]
                        I_eff = I_end - ss
                        if I_eff > 0:
                            c = I_eff * h_cost
                            hold += c
                            tag = f"H={c:.2f}"
                        elif I_eff < 0:
                            c = (-I_eff) * p_cost
                            slack += c
                            tag = f"S={c:.2f}"
                        else:
                            tag = ""
                        log_list.append(f"   {W+1:>4} | {I:>6.2f} | {PR[W]:>6.2f} | {GR[W]:>6.2f} | {I_end:>6.2f} | {tag}")
                        I = I_end
                    C = hold + slack
                    cost_set.append(C)
                    log_list.append(f"   → Scenario cost = {C:.2f}  (hold={hold:.2f}, slack={slack:.2f})")
                return POW_vec, PR_nom, cost_set

            # ------- Evaluate all policies -------
            policy_cost_sets = []
            policy_pow_nom_cache = []  # (policy, POW_vec, PR_nom)
            all_policy_logs = []  # store detailed logs; then flush to main

            for pol_idx, policy in enumerate(policies, start=1):
                plogs = []
                POW_vec, PR_nom, cost_set = cost_for_policy(policy, plogs)
                policy_pow_nom_cache.append((policy, POW_vec, PR_nom))
                policy_cost_sets.append([-c for c in cost_set])  # NEGATE for comparator
                # Full logs for this policy (like Simple)
                all_policy_logs.extend(plogs)
                # Compact per-policy summary line
                logs.append(f"  Policy {pol_idx}/{len(policies)}  LT_orders={policy}  → grey cost set size={len(cost_set)}; sample={cost_set[:min(4,len(cost_set))]}")

            # Flush detailed per-policy logs after summaries
            logs.extend(all_policy_logs)

            # ------- Compare policies with discrete grey numbers -------
            comp_items = [(f"LT_orders{tuple(policy)}", ",".join(str(v) for v in vals), "") for policy, vals in zip([p for p,_,_ in policy_pow_nom_cache], policy_cost_sets)]
            if len(comp_items) == 1:
                labels = [comp_items[0][0]]
                scores = [0.0]
                order_desc = [0]
            else:
                try:
                    labels, _Mgt, _Meq, _Madv, scores, order_desc, _log_text = DiscreteGreyNumberComparator.compare_many(comp_items)
                except Exception as e:
                    messagebox.showerror("Comparison error", f"Grey discrete comparison failed for item {it}:\n{e}")
                    return

            # Print ONLY final scores (no matrices)
            logs.append("\n  Grey score summary (higher is better because we used negated costs):")
            for lbl, sc in zip(labels, scores):
                logs.append(f"    {lbl}: score={sc:.6f}")

            best_idx = order_desc[0]
            best_policy = policies[best_idx]
            logs.append(f"\n  Chosen complex policy for {it}: LT_orders per demand week = {best_policy} (weeks { [w+1 for w in demand_weeks] })")

            # Fix POW & PR_nom to chosen policy's nominal plan
            best_pow = [0.0]*H
            best_prn = [0.0]*H
            for pol, pow_vec, pr_nom in policy_pow_nom_cache:
                if pol == best_policy:
                    best_pow = pow_vec; best_prn = pr_nom; break
            POW[it] = best_pow
            PRNOM[it] = best_prn
            lt_policy_choice[it] = {demand_weeks[i]: best_policy[i] for i in range(len(demand_weeks))}

        # Final tables: PR_nominal and POW
        logs.append("\n" + "="*72)
        logs.append("FINAL Planned Order Receipts (PR_nominal):")
        head = "Item \\ Week  " + "  ".join(f"{w:>6}" for w in self.model['weeks'])
        logs.append(head)
        for it in order:
            row = f"{it:>11}  " + "  ".join(f"{x:>6.2f}" for x in PRNOM[it])
            logs.append(row)

        logs.append("\n" + "="*72)
        logs.append("FINAL Planned Order Releases (POW):")
        head = "Item \\ Week  " + "  ".join(f"{w:>6}" for w in self.model['weeks'])
        logs.append(head)
        for it in order:
            row = f"{it:>11}  " + "  ".join(f"{x:>6.2f}" for x in POW[it])
            logs.append(row)

        # Note about horizon vs releases
        logs.append("\nNOTE: If horizon starts at W9 and LT=1, the release for W9 occurs at W8 (outside the table).")
        logs.append("      That's why POW may show a single in-horizon release while PR_nominal shows receipts at W9 & W10.")

        self.results = {"pow": POW, "pr_nominal": PRNOM, "lt_policy": lt_policy_choice, "order": order, "weeks": self.model['weeks']}
        self.result_widget.delete("1.0", tk.END)
        self.result_widget.insert(tk.END, "\n".join(logs) + "\n")

    # ---------------------- Helpers ----------------------
    def _build_gr(self, it, H, mps, POW, bom):
        # GR = MPS self + sum(parent POW * usage)
        GR = mps[it]['gr'][:] if it in mps else [0.0]*H
        for (p, c, u) in bom:
            if c == it:
                parent_pow = POW[p]
                for t in range(H):
                    GR[t] += u * parent_pow[t]
        return GR

    def _read_excel_model(self, path):
        wb = load_workbook(path, data_only=True)

        def find_sheet(names):
            for s in wb.sheetnames:
                low = s.strip().lower()
                for k in names:
                    if k in low:
                        return wb[s]
            return None

        ws_items = find_sheet(["item"])
        ws_bom = find_sheet(["bom"])
        ws_mps = find_sheet(["mps"]) or find_sheet(["mps & costs","mps&costs","cost"])

        if ws_items is None or ws_bom is None or ws_mps is None:
            raise ValueError("Cannot find all required sheets. Need: Items, BOM, MPS & Costs.")

        # Items
        items = {}
        rows = list(ws_items.iter_rows(values_only=True))
        if not rows:
            raise ValueError("Items sheet is empty.")
        header = [str(x).strip().lower() if x is not None else "" for x in rows[0]]
        def idx(colname):
            for i, h in enumerate(header):
                if colname in h:
                    return i
            return None
        i_item = idx("item")
        i_ll = idx("lower")
        i_ul = idx("upper")
        i_lot = idx("lot")
        i_ss = idx("safety")
        i_inv = idx("invent")
        for r in rows[1:]:
            if r is None: continue
            name = (r[i_item] or "").strip() if (i_item is not None and r[i_item] is not None) else ""
            if not name: continue
            items[name] = {
                "lt_low": int(r[i_ll]) if (i_ll is not None and r[i_ll] is not None) else 1,
                "lt_up":  int(r[i_ul]) if (i_ul is not None and r[i_ul] is not None) else 1,
                "lot": float(r[i_lot]) if (i_lot is not None and r[i_lot] is not None) else 0.0,
                "ss":  float(r[i_ss])  if (i_ss  is not None and r[i_ss]  is not None) else 0.0,
                "I0":  float(r[i_inv]) if (i_inv is not None and r[i_inv] is not None) else 0.0,
            }

        # BOM
        bom = []
        rows = list(ws_bom.iter_rows(values_only=True))
        if rows:
            h = [str(x).strip().lower() if x is not None else "" for x in rows[0]]
            i_p = next((i for i,v in enumerate(h) if "parent" in v), 0)
            i_c = next((i for i,v in enumerate(h) if "child" in v), 1)
            i_u = next((i for i,v in enumerate(h) if "usage" in v), 2)
            for r in rows[1:]:
                if r is None: continue
                p = (r[i_p] or "").strip() if r[i_p] is not None else ""
                c = (r[i_c] or "").strip() if r[i_c] is not None else ""
                if not p or not c: continue
                u = float(r[i_u]) if r[i_u] is not None else 1.0
                bom.append((p,c,u))

        # MPS & Costs
        mps = {}
        rows = list(ws_mps.iter_rows(values_only=True))
        if not rows:
            raise ValueError("MPS & Costs sheet is empty.")
        header = [str(x).strip() if x is not None else "" for x in rows[0]]
        week_cols = []
        for j,h in enumerate(header):
            if h and h.strip().lower().startswith("w"):
                week_cols.append(j)
        if not week_cols:
            raise ValueError("No week columns (W1, W2, ...) found in MPS & Costs.")
        weeks = [header[j] for j in week_cols]
        i_item = next((i for i,v in enumerate(header) if "item" in v.lower()), None)
        i_h = next((i for i,v in enumerate(header) if "holding" in v.lower()), None)
        i_p = next((i for i,v in enumerate(header) if "slack" in v.lower() or "shortage" in v.lower() or "penalty" in v.lower()), None)

        for r in rows[1:]:
            if r is None: continue
            name = (r[i_item] or "").strip() if (i_item is not None and r[i_item] is not None) else ""
            if not name: continue
            h = float(r[i_h]) if (i_h is not None and r[i_h] is not None) else 0.0
            p = float(r[i_p]) if (i_p is not None and r[i_p] is not None) else 0.0
            gr = [float(r[j] or 0.0) for j in week_cols]
            mps[name] = {"h":h, "p":p, "gr":gr}

        return {"items":items, "bom":bom, "mps":mps, "weeks":weeks}

    # ---------------------- Utils ----------------------
    def _topo_order(self, items, bom):
        items = list(items)
        indeg = {x:0 for x in items}
        children = defaultdict(list)
        for p,c,u in bom:
            if p not in indeg: indeg[p]=0
            if c not in indeg: indeg[c]=0
            indeg[c] += 1
            children[p].append(c)
        q = deque([x for x in indeg if indeg[x]==0])
        order = []
        while q:
            u = q.popleft()
            if u in order:
                continue
            if u in items:
                order.append(u)
            for v in children.get(u, []):
                indeg[v] -= 1
                if indeg[v]==0:
                    q.append(v)
        for x in items:
            if x not in order:
                order.append(x)
        return order

    # ---------------------- Save/Load project ----------------------
    def save_project(self):
        data = {"excel_path": self.file_var.get().strip()}
        path = filedialog.asksaveasfilename(defaultextension=".json", filetypes=[("JSON","*.json")], title="Save Grey MRP Complex Project")
        if not path: return
        try:
            with open(path, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=2, ensure_ascii=False)
            messagebox.showinfo("Saved", f"Saved to:\n{path}")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to save:\n{e}")

    def load_project(self):
        path = filedialog.askopenfilename(defaultextension=".json", filetypes=[("JSON","*.json")], title="Load Grey MRP Complex Project")
        if not path: return
        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
        except Exception as e:
            messagebox.showerror("Error", f"Failed to load:\n{e}")
            return
        self.file_var.set(data.get("excel_path",""))
        self.model = None
        self.results = None
        self.result_widget.delete("1.0", tk.END)
        self.result_widget.insert(tk.END, "Project loaded. Use Preview or Run.\n")


# Optional manual test
if __name__ == "__main__":
    root = tk.Tk()
    root.title("Grey MRP — Complex Approach")
    text = tk.Text(root, bg="#111", fg="#eee", width=120, height=40)
    text.pack(fill="both", expand=True)
    frame = GreyMRPComplexInput(root, text)
    frame.pack(fill="x")
    root.mainloop()
